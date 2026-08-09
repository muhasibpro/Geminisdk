import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_tovarlar_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tovarlar"

    # Turn on gridlines
    ws.views.sheetView[0].showGridLines = True

    # Header style
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Data styles
    font_regular = Font(name="Calibri", size=11)
    font_bold = Font(name="Calibri", size=11, bold=True)
    
    fill_even = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    thin_border_side = Side(border_style="thin", color="D9D9D9")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    thick_bottom_side = Side(border_style="medium", color="1F4E79")
    double_bottom_side = Side(border_style="double", color="1F4E79")
    top_thin_side = Side(border_style="thin", color="1F4E79")

    # Title Banner
    ws.merge_cells("A1:D1")
    title_cell = ws["A1"]
    title_cell.value = "KIRIM TOVARLARI RO'YXATI"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    # Table Headers
    headers = ["№", "Tovar nomi", "Shtrixkod (Barcode)", "Kirim narxi (so'm)"]
    ws.append([]) # Row 2 empty space
    
    ws.row_dimensions[3].height = 26
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thick_bottom_side)

    # Data
    products = [
        ("Coca-Cola 1.5L", "4780000123456", 11500),
        ("Pepsi 1.5L", "4780000123463", 11000),
        ("Chortoq 0.5L (Gazli)", "4780000234567", 4500),
        ("Nestle Sut 1L 3.2%", "4780000345678", 12000),
        ("Lays Chipslari 140g", "4780000456789", 14500),
        ("Crafers Shokolad 100g", "4780000567890", 9000),
        ("Musaffo Sariyog' 200g", "4780000678901", 22000),
        ("MacCoffee 3in1 (1 dona)", "4780000789012", 2200),
        ("Lipton Choy 100g", "4780000890123", 16500),
        ("Dobriy Sharbat 1L", "4780000901234", 13000),
    ]

    start_row = 4
    for idx, (name, barcode, price) in enumerate(products, 1):
        row_num = start_row + idx - 1
        ws.row_dimensions[row_num].height = 22
        
        c1 = ws.cell(row=row_num, column=1, value=idx)
        c2 = ws.cell(row=row_num, column=2, value=name)
        c3 = ws.cell(row=row_num, column=3, value=str(barcode))
        c4 = ws.cell(row=row_num, column=4, value=price)

        # Alignments
        c1.alignment = Alignment(horizontal="center", vertical="center")
        c2.alignment = Alignment(horizontal="left", vertical="center")
        c3.alignment = Alignment(horizontal="center", vertical="center")
        c4.alignment = Alignment(horizontal="right", vertical="center")

        # Number format for prices
        c4.number_format = '#,##0 "so\'m"'

        # Row Fill & Borders
        fill = fill_even if idx % 2 == 0 else fill_odd
        for cell in (c1, c2, c3, c4):
            cell.font = font_regular
            cell.fill = fill
            cell.border = thin_border

    # Total Row
    total_row = start_row + len(products)
    ws.row_dimensions[total_row].height = 25
    
    ws.cell(row=total_row, column=1, value="")
    total_label = ws.cell(row=total_row, column=2, value="JAMI / UMUMIY SUMMA:")
    total_label.font = font_bold
    total_label.alignment = Alignment(horizontal="right", vertical="center")
    
    ws.cell(row=total_row, column=3, value="")
    
    total_val = ws.cell(row=total_row, column=4, value=f"=SUM(D{start_row}:D{total_row-1})")
    total_val.font = font_bold
    total_val.alignment = Alignment(horizontal="right", vertical="center")
    total_val.number_format = '#,##0 "so\'m"'

    summary_border = Border(top=top_thin_side, bottom=double_bottom_side, left=thin_border_side, right=thin_border_side)
    for col in range(1, 5):
        ws.cell(row=total_row, column=col).border = summary_border

    # Adjust Column Widths dynamically
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Skip merged banner title row for length calc
            if cell.row == 1:
                continue
            val_str = str(cell.value or "")
            if cell.number_format and "so'm" in cell.number_format and isinstance(cell.value, (int, float)):
                val_str = f"{cell.value:,.0f} so'm"
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

    output_filename = "tovarlar.xlsx"
    wb.save(output_filename)
    print(f"File successfully created: {output_filename}")

if __name__ == "__main__":
    generate_tovarlar_excel()
